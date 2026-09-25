#!/usr/bin/env python3
"""Export every price on the site to one workbook — the central pricing sheet.

    python3 scripts/export-pricing-sheet.py               # → ~/Downloads/Product Images/_data/AC_Supply_Pricing.xlsx
    python3 scripts/export-pricing-sheet.py out.xlsx      # somewhere else (never inside the repo)
    python3 scripts/export-pricing-sheet.py --compare a2e8084   # add Was / Change columns and a Changes tab vs that commit

Generated, one-way: data/products/*.json stays the single editable store. Change a
price there, run build-products.py, then re-run this. The workbook carries the
private companion data (sources, supplier costs, margins), so it is written
OUTSIDE the repo — the repo root is what Cloudflare serves.

Tabs: Read me · Summary (one row per product) · Prices (every variant × tier, exactly as
the site shows it) · Costs & margins (supplier-costed ladders from data/private/*.json).
With --compare REF: Prices gains Was £ / Change / vs-REF columns and a Changes tab lists every
price changed, added or removed since the data/products/*.json at that git ref.
"""
import importlib.util
import json
import re
import subprocess
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from productlib import IMAGE_SOURCE, ROOT, SITE, BuildError, load_records

DEFAULT_OUT = IMAGE_SOURCE / '_data' / 'AC_Supply_Pricing.xlsx'

# Reuse the build's own tier maths so the sheet can never disagree with the pages.
_spec = importlib.util.spec_from_file_location('build_products', Path(__file__).with_name('build-products.py'))
build = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(build)

FONT = 'Arial'
HEAD = Font(name=FONT, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', start_color='0F1F3D')  # --ink
BODY = Font(name=FONT)
INPUT = Font(name=FONT, color='0000FF')  # hardcoded data from the JSON records
BOLD = Font(name=FONT, bold=True)
UNIT_FMT = '£#,##0.000'
MONEY_FMT = '£#,##0.00'
PCT_FMT = '0.0%;-0.0%;-'


def page_url(record):
    if record['status'] == 'published':
        return SITE + record['seo']['canonical']
    if record['status'] == 'preview':
        return f'{SITE}/preview/{record["slug"]}/'
    return ''


def variant_label(record, selection):
    labels = []
    for option in record['options']:
        value = next((v for v in option['values'] if v['value'] == selection.get(option['key'])), None)
        if value:
            labels.append(value['label'])
    return ' · '.join(labels)


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for col, width in enumerate(widths, 1):
        cell = ws.cell(row=1, column=col)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'
    return ws


def style_row(ws, row, fonts):
    for col, font in enumerate(fonts, 1):
        ws.cell(row=row, column=col).font = font


def cost_rows(record):
    """[(ref, qty, trade cost, basis)] from the private supplier costings; [] when there are none."""
    costing = record.get('internal', {}).get('express_costing')
    if not costing:
        return []
    supplier = costing.get('supplier', '')
    rows = []
    if 'ladders' in costing:  # cost already worked out per ref and qty
        for ref, ladder in costing['ladders'].items():
            for rung in ladder:
                if 'price' in rung:
                    rows.append((ref, rung['qty'], Decimal(rung['trade_cost']), supplier))
    elif 'list_prices_gbp_ex_vat' in costing:  # list-price ladders: cost worked out at the rung at or below qty
        discount = Decimal(costing['trade_discount_off_list'])
        refs = {p['ref']: p for p in record['pricing']}
        for name, ladder in costing['list_prices_gbp_ex_vat'].items():
            m = re.fullmatch(r'(\d+)oz (\d) colours?', name)
            ref = m and f'clc-{m[1]}oz-express-{m[2]}c'
            if ref not in refs:
                continue
            rungs = {int(q): Decimal(u) for q, u in ladder.items()}
            for tier in refs[ref]['tiers']:
                below = [q for q in rungs if q <= tier['qty']]
                if below:
                    rows.append((ref, tier['qty'], rungs[max(below)] * (1 - discount), f'{supplier} — list at {max(below):,} less {discount:.0%}'))
    return rows


def old_prices(ref):
    """{(slug, sku, qty): (unit, product title, variant label)} as the site showed them at git ref."""
    git = lambda *a: subprocess.run(['git', '-C', str(ROOT), *a], capture_output=True, text=True, check=True).stdout
    out = {}
    for path in git('ls-tree', '--name-only', ref, 'data/products/').split():
        record = json.loads(git('show', f'{ref}:{path}'))
        for variant in record['variants']:
            for qty, unit in build.priced_tiers(record, variant, variant['selection']):
                if unit is not None:
                    out[(record['slug'], variant['sku'], qty)] = (unit, record['title'], variant_label(record, variant['selection']))
    return out


def main():
    args = sys.argv[1:]
    compare = None
    if '--compare' in args:
        i = args.index('--compare'); compare = args[i + 1]; del args[i:i + 2]
    out = Path(args[0]).expanduser().resolve() if args else DEFAULT_OUT
    if ROOT.resolve() in out.parents:
        raise BuildError(f'{out} is inside the repo, which is publicly served — this sheet holds costs and margins; write it elsewhere')
    records = [r for _, r in load_records()]
    was = old_prices(compare) if compare else {}
    if compare:
        label = subprocess.run(['git', '-C', str(ROOT), 'log', '-1', '--format=%cd', '--date=format:%d %b %Y', compare], capture_output=True, text=True).stdout.strip()
        vs = f'vs {label}'
    seen, poa = set(), set()

    wb = Workbook()
    readme = wb.active
    readme.title = 'Read me'

    # ---- Prices: one row per variant × tier ----
    prices = sheet(wb, 'Prices', ['Product', 'Slug', 'Status', 'SKU', 'Variant', 'Availability', 'Qty', 'Unit £ (ex VAT)', 'Order total £ (ex VAT)',
                                  'Pricing ref', 'Variant count', 'Page'] + (['Was £ (ex VAT)', 'Change', vs] if compare else []),
                   [30, 22, 11, 26, 44, 13, 10, 14, 16, 30, 8, 52] + ([14, 10, 12] if compare else []))
    row = 1
    for record in records:
        url = page_url(record)
        for variant in record['variants']:
            selection = variant['selection']
            tiers = build.priced_tiers(record, variant, selection)
            if tiers:
                availability = 'Priced'
            elif variant['available'] or build.variant_mode(record, variant) == 'quote_only':
                availability = 'Quote only'
            else:
                availability = 'Not offered'
            for i, (qty, unit) in enumerate(tiers or [(None, None)]):
                row += 1
                prices.append([record['title'], record['slug'], record['status'], variant['sku'], variant_label(record, selection), availability,
                               qty, float(unit) if unit is not None else ('POA' if qty else None),
                               f'=IF(ISNUMBER(H{row}),ROUND(G{row}*H{row},2),"")',
                               variant.get('pricing_ref', ''), 1 if i == 0 else 0, url])
                if compare:
                    old = was.get((record['slug'], variant['sku'], qty)) if unit is not None else None
                    (seen if unit is not None else poa).add((record['slug'], variant['sku'], qty))
                    state = '' if unit is None else 'New' if old is None else 'Same' if old[0] == unit else 'Up' if unit > old[0] else 'Down'
                    prices.cell(row=row, column=13, value=float(old[0]) if old else None).number_format = UNIT_FMT
                    prices.cell(row=row, column=14, value=f'=IF(AND(ISNUMBER(H{row}),ISNUMBER(M{row})),H{row}/M{row}-1,"")').number_format = PCT_FMT
                    prices.cell(row=row, column=15, value=state)
                    for col in (13, 14, 15):
                        prices.cell(row=row, column=col).font = INPUT if col == 13 else BODY
                style_row(prices, row, [BODY] * 6 + [INPUT, INPUT, BODY, BODY, BODY, BODY])
                prices.cell(row=row, column=7).number_format = '#,##0'
                prices.cell(row=row, column=8).number_format = UNIT_FMT
                prices.cell(row=row, column=9).number_format = MONEY_FMT
    prices.auto_filter.ref = f'A1:{"O" if compare else "L"}{row}'
    prices.column_dimensions['K'].hidden = True  # helper: 1 on each variant's first row, counted by Summary
    last = row

    # ---- Summary: one row per product ----
    summary = sheet(wb, 'Summary', ['Product', 'Slug', 'Category', 'Status', 'Variants', 'Priced', 'Quote only', 'Not offered',
                                    'Lowest unit £', 'Highest unit £', 'Tier quantities', 'Pricing source', 'Review date', 'Page'],
                    [30, 22, 20, 11, 9, 8, 10, 11, 13, 13, 26, 60, 12, 52])
    rng = lambda col: f'Prices!${col}$2:${col}${last}'
    for r, record in enumerate(records, 2):
        tiers = sorted({t['qty'] for p in record['pricing'] for t in p['tiers']})
        internal = record.get('internal', {})
        summary.append([record['title'], record['slug'], record['category'], record['status'],
                        f'=COUNTIFS({rng("B")},B{r},{rng("K")},1)',
                        f'=COUNTIFS({rng("B")},B{r},{rng("F")},"Priced",{rng("K")},1)',
                        f'=COUNTIFS({rng("B")},B{r},{rng("F")},"Quote only")',
                        f'=COUNTIFS({rng("B")},B{r},{rng("F")},"Not offered")',
                        f'=IF(F{r}=0,"",_xlfn.MINIFS({rng("H")},{rng("B")},B{r}))',
                        f'=IF(F{r}=0,"",_xlfn.MAXIFS({rng("H")},{rng("B")},B{r}))',
                        ' / '.join(f'{q:,}' for q in tiers),
                        internal.get('pricing_source', '(no private file)'), internal.get('review_date', ''), page_url(record)])
        style_row(summary, r, [BOLD, BODY, BODY, BODY] + [BODY] * 6 + [INPUT, INPUT, INPUT, BODY])
        for col in (9, 10):
            summary.cell(row=r, column=col).number_format = UNIT_FMT
    summary.auto_filter.ref = f'A1:N{len(records) + 1}'

    # ---- Costs & margins: supplier-costed ladders only ----
    costs = sheet(wb, 'Costs & margins', ['Product', 'Pricing ref', 'Qty', 'Trade cost £/unit', 'Sell £/unit', 'Margin on sale', 'Target margin', 'Check', 'Cost basis'],
                  [30, 32, 10, 15, 13, 13, 12, 14, 70])
    row = 1
    for record in records:
        target = record.get('internal', {}).get('express_costing', {}).get('target_margin_on_sale')
        ladders = {p['ref']: {t['qty']: t['unit'] for t in p['tiers']} for p in record['pricing']}
        for ref, qty, cost, basis in cost_rows(record):
            sell = ladders.get(ref, {}).get(qty)
            if sell is None:
                continue
            row += 1
            costs.append([record['title'], ref, qty, float(cost), float(Decimal(sell)),
                          f'=IF(E{row}=0,"",(E{row}-D{row})/E{row})', float(Decimal(target)) if target else None,
                          f'=IF(OR(G{row}="",F{row}=""),"",IF(F{row}<G{row},"BELOW TARGET","OK"))', basis])
            style_row(costs, row, [BODY, BODY, INPUT, INPUT, INPUT, BODY, INPUT, BODY, BODY])
            costs.cell(row=row, column=3).number_format = '#,##0'
            costs.cell(row=row, column=4).number_format = '£#,##0.0000'
            costs.cell(row=row, column=5).number_format = UNIT_FMT
            costs.cell(row=row, column=6).number_format = PCT_FMT
            costs.cell(row=row, column=7).number_format = PCT_FMT
    costs.auto_filter.ref = f'A1:I{max(row, 2)}'
    costs['D1'].comment = Comment('Trade cost per unit — from data/private/<slug>.json (internal.express_costing).', 'export')
    costs['G1'].comment = Comment('Minimum margin on sale — from data/private/<slug>.json (internal.express_costing).', 'export')

    # ---- Changes: every price changed, added or removed since --compare ----
    if compare:
        ch = wb.create_sheet('Changes')
        lines = []
        for r in range(2, last + 1):
            state = prices.cell(row=r, column=15).value
            if state in ('Up', 'Down', 'New'):
                lines.append((prices.cell(row=r, column=1).value, prices.cell(row=r, column=4).value, prices.cell(row=r, column=5).value,
                              prices.cell(row=r, column=7).value, prices.cell(row=r, column=13).value, prices.cell(row=r, column=8).value, state))
        for (slug, sku, qty), (unit, title, var) in was.items():
            if (slug, sku, qty) not in seen:
                lines.append((title, sku, var, qty, float(unit), None, 'Now POA' if (slug, sku, qty) in poa else 'Removed'))
        order = {'Up': 0, 'Down': 1, 'New': 2, 'Removed': 3, 'Now POA': 4}
        first, end = 5, len(lines) + 4
        for i, state in enumerate(order):  # counts above the table
            ch.cell(row=1, column=1 + i * 2, value=state).font = BOLD
            ch.cell(row=1, column=2 + i * 2, value=f'=COUNTIF($H${first}:$H${end},"{state}")').font = BODY
        ch.cell(row=2, column=1, value=f'Every site price changed, added or removed {vs} (git {compare}). Removed = that option × quantity is no longer offered; Now POA = still listed, price on request.').font = BODY
        for c, (head, width) in enumerate(zip(['Product', 'SKU', 'Variant', 'Qty', 'Was £', 'Now £', 'Change', 'What happened'], [30, 28, 48, 10, 12, 12, 10, 14]), 1):
            cell = ch.cell(row=4, column=c, value=head); cell.font, cell.fill = HEAD, HEAD_FILL
            ch.column_dimensions[get_column_letter(c)].width = width
        for n, (title, sku, var, qty, old, new, state) in enumerate(sorted(lines, key=lambda l: (l[0], order[l[6]], l[1], l[3])), first):
            for c, v in enumerate([title, sku, var, qty, old, new, f'=IF(AND(ISNUMBER(E{n}),ISNUMBER(F{n})),F{n}/E{n}-1,"")', state], 1):
                ch.cell(row=n, column=c, value=v)
            style_row(ch, n, [BODY, BODY, BODY, INPUT, INPUT, INPUT, BODY, BODY])
            ch.cell(row=n, column=4).number_format = '#,##0'
            for c in (5, 6): ch.cell(row=n, column=c).number_format = UNIT_FMT
            ch.cell(row=n, column=7).number_format = PCT_FMT
        ch.freeze_panes = f'A{first}'
        ch.auto_filter.ref = f'A4:H{end}'

    # ---- Read me ----
    notes = [
        ('AC Supply Co — central pricing sheet', BOLD),
        (f'Generated {date.today():%d %B %Y} by scripts/export-pricing-sheet.py from data/products/*.json + data/private/*.json.', BODY),
        ('', BODY),
        ('Do not edit prices here — this file is overwritten on every export.', BOLD),
        ('To change a price: edit data/products/<slug>.json → python3 scripts/build-products.py → re-run this export.', BODY),
        ('', BODY),
        ('Prices: every variant × quantity tier exactly as the site shows it (ex VAT). Order total = qty × unit, rounded to the penny like the site.', BODY),
        ('Summary: one row per product with counts, price range, tiers, pricing source and review date.', BODY),
        ('Costs & margins: every ladder priced from a supplier cost sheet. Margin = (sell − cost) ÷ sell.', BODY),
        ('Blue text = values copied from the records; black = formulas.', BODY),
        *([(f'Changes: every price changed, added or removed {vs}. Prices also shows the old price (Was £) and the % change on each row.', BODY)] if compare else []),
        ('', BODY),
        ('PRIVATE: contains supplier costs and margins. Never put this file in the website repo or send it to customers.', BOLD),
    ]
    for text, font in notes:
        readme.append([text])
        readme.cell(row=readme.max_row, column=1).font = font
    readme.column_dimensions['A'].width = 120

    wb.move_sheet('Summary', offset=-1)
    wb.calculation.fullCalcOnLoad = True  # openpyxl stores no cached values; make Excel/Sheets/Numbers compute on open
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f'wrote {out} — {len(records)} products, {last - 1} price rows, {row - 1} costed rows')


if __name__ == '__main__':
    try:
        main()
    except BuildError as err:
        sys.exit(f'error: {err}')
